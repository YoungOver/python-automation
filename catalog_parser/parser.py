"""Каталог-парсер: books.toscrape.com -> Excel-отчёт + сводка JSON.

Запуск:  python parser.py --out books.xlsx --concurrency 8
Сайт books.toscrape.com создан специально для тренировки парсинга.
"""
import argparse
import asyncio
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import httpx
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = 'https://books.toscrape.com/'
RATING = {'One': 1, 'Two': 2, 'Three': 3, 'Four': 4, 'Five': 5}
GBP_RUB = 118.0  # курс для пересчёта в рубли


async def get(client, url, sem, tries=4):
    for attempt in range(tries):
        try:
            async with sem:
                r = await client.get(url, timeout=30)
            r.raise_for_status()
            return r.text
        except (httpx.HTTPError, httpx.TimeoutException):
            await asyncio.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f'не удалось загрузить {url}')


def parse_list(html, page_url):
    soup = BeautifulSoup(html, 'html.parser')
    links = [urljoin(page_url, a['href']) for a in soup.select('article.product_pod h3 a')]
    nxt = soup.select_one('li.next a')
    return links, (urljoin(page_url, nxt['href']) if nxt else None)


def parse_book(html, url):
    s = BeautifulSoup(html, 'html.parser')
    table = {tr.th.text.strip(): tr.td.text.strip() for tr in s.select('table.table tr')}
    price = float(re.sub(r'[^\d.]', '', table.get('Price (incl. tax)', '0')))
    stock = int((re.search(r'(\d+)', table.get('Availability', '')) or [0, 0])[1])
    rating_cls = s.select_one('p.star-rating')['class'][1]
    return {
        'title': s.select_one('div.product_main h1').text.strip(),
        'category': s.select('ul.breadcrumb li a')[-1].text.strip(),
        'price_gbp': price,
        'price_rub': round(price * GBP_RUB),
        'rating': RATING.get(rating_cls, 0),
        'stock': stock,
        'upc': table.get('UPC', ''),
        'url': url,
    }


async def crawl(concurrency):
    sem = asyncio.Semaphore(concurrency)
    headers = {'User-Agent': 'CatalogParser/1.0 (+portfolio demo)'}
    async with httpx.AsyncClient(headers=headers, follow_redirects=True) as client:
        links, page = [], urljoin(BASE, 'catalogue/page-1.html')
        while page:
            found, page = parse_list(await get(client, page, sem), page)
            links += found
        pages = await asyncio.gather(*(get(client, u, sem) for u in links))
        return [parse_book(h, u) for h, u in zip(pages, links)]


def write_excel(books, path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Товары'
    cols = [('Название', 'title', 48), ('Категория', 'category', 20), ('Цена, £', 'price_gbp', 10),
            ('Цена, ₽', 'price_rub', 11), ('Рейтинг', 'rating', 9), ('Остаток, шт', 'stock', 12),
            ('UPC', 'upc', 18), ('Ссылка', 'url', 60)]
    ws.append([c[0] for c in cols])
    for b in sorted(books, key=lambda x: (x['category'], -x['rating'])):
        ws.append([b[c[1]] for c in cols])
    for i, (_, _, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    n = len(books) + 1
    tab = Table(displayName='Books', ref=f'A1:H{n}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = 'A2'
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = '# ##0 ₽'
    ws.conditional_formatting.add(f'E2:E{n}', ColorScaleRule(start_type='num', start_value=1, start_color='F8696B',
                                                             mid_type='num', mid_value=3, mid_color='FFEB84',
                                                             end_type='num', end_value=5, end_color='63BE7B'))

    agg = defaultdict(lambda: {'n': 0, 'sum': 0.0, 'r': 0, 'stock': 0})
    for b in books:
        a = agg[b['category']]
        a['n'] += 1; a['sum'] += b['price_rub']; a['r'] += b['rating']; a['stock'] += b['stock']
    summary = sorted(({'category': k, 'count': v['n'], 'avg_rub': round(v['sum'] / v['n']),
                       'avg_rating': round(v['r'] / v['n'], 2), 'stock': v['stock']} for k, v in agg.items()),
                     key=lambda x: -x['count'])
    ws2 = wb.create_sheet('Сводка')
    ws2.append(['Категория', 'Товаров', 'Средняя цена, ₽', 'Средний рейтинг', 'Остаток, шт'])
    for s in summary:
        ws2.append([s['category'], s['count'], s['avg_rub'], s['avg_rating'], s['stock']])
    for c, w in zip('ABCDE', (24, 10, 16, 16, 12)):
        ws2.column_dimensions[c].width = w
    for cell in ws2[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center')
    chart = BarChart(); chart.type = 'bar'; chart.style = 10
    chart.title = 'Топ-15 категорий по числу товаров'; chart.height = 11; chart.width = 20
    top = min(15, len(summary)) + 1
    chart.add_data(Reference(ws2, min_col=2, min_row=1, max_row=top), titles_from_data=True)
    chart.set_categories(Reference(ws2, min_col=1, min_row=2, max_row=top))
    ws2.add_chart(chart, 'G2')
    wb.save(path)
    return summary


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='books.xlsx')
    ap.add_argument('--concurrency', type=int, default=8)
    args = ap.parse_args()
    t0 = datetime.now()
    books = asyncio.run(crawl(args.concurrency))
    summary = write_excel(books, args.out)
    secs = (datetime.now() - t0).total_seconds()
    meta = {'total': len(books), 'categories': len(summary), 'seconds': round(secs, 1),
            'avg_rub': round(sum(b['price_rub'] for b in books) / len(books)),
            'in_stock': sum(b['stock'] for b in books), 'summary': summary[:12],
            'top': sorted(books, key=lambda b: (-b['rating'], -b['price_rub']))[:6],
            'generated': datetime.now().strftime('%d.%m.%Y %H:%M')}
    Path(args.out).with_suffix('.json').write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f"{len(books)} товаров, {len(summary)} категорий за {secs:.0f} с -> {args.out}")


if __name__ == '__main__':
    main()
